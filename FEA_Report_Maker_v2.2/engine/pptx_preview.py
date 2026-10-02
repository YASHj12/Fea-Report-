#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
pptx_preview.py  -  draw the slides of a .pptx as pictures, with nothing but Pillow.

No PowerPoint, no LibreOffice, no internet.  It reads the finished file (so the preview is exactly what is inside the
download) and draws what mw_report.py writes: layout pictures and rectangles, the title placeholder, text boxes (wrap,
alignment, bullets, numbering, line pitch, paragraph spacing), pictures, tables (fills, white borders, auto-grown rows)
and straight lines.  Fonts: the bundled Liberation Sans (metric-compatible with Arial).

    deck = Deck("report.pptx")
    deck.count()                       -> number of slides
    deck.title(0)                      -> "Static Structural Analysis Of ..."
    deck.render(3, width_px=1280)      -> PIL.Image  (RGB)
    deck.text_lines(3)                 -> [(x, baseline_y, text, size_pt, bold, is_label)]   (points; used by the self-test)
    deck.unsupported                   -> shapes that were skipped (should stay empty for decks made by mw_report.py)
    images_to_pdf(deck, "out.pdf")     -> a picture-based PDF (the fallback when LibreOffice is not installed)

    python pptx_preview.py deck.pptx outdir [width]      writes outdir/s01.png ...
"""
from __future__ import annotations

import glob
import hashlib
import io
import sys
import threading
from pathlib import Path

from lxml import etree
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
PN = "http://schemas.openxmlformats.org/presentationml/2006/main"
RN = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS = {"a": A, "p": PN, "r": RN}
EMU = 12700.0
HERE = Path(__file__).resolve().parent

# Line metrics (fractions of the font size), measured on LibreOffice's own PDF output (tests/test_preview.py):
#   single spacing  -> line pitch 1.200 x size, first baseline 1.000 x size below the top of the text area
#   exact spacing   -> first baseline = pitch - 0.2 x size
# (PowerPoint's single spacing for Arial is also 1.2 x size, which is why the template's 14 pt notes are 16.8 pt apart.)
DESC = 0.2
SINGLE = 1.2
FIT = 1.001                          # line-fitting tolerance: my widths are 0.08 % above LibreOffice's on average (measured on 245 text runs)
DEFAULT_SIZE = 18.0
DEFAULT_COLOR = "23282D"


# ───────────────────────────────────────────────── fonts ─────────────────────────────────────────────────
_FONT_FILES = {(False, False): "LiberationSans-Regular.ttf", (True, False): "LiberationSans-Bold.ttf",
               (False, True): "LiberationSans-Italic.ttf", (True, True): "LiberationSans-BoldItalic.ttf"}
_FONT_PATH: dict = {}
_FONTS: dict = {}
_FONT_LOCK = threading.Lock()


def _font_path(bold: bool, italic: bool):
    key = (bool(bold), bool(italic))
    if key in _FONT_PATH:
        return _FONT_PATH[key]
    cands = []
    for k in (key, (key[0], False), (False, False)):              # italic missing -> upright, bold missing -> regular
        name = _FONT_FILES[k]
        cands += [str(HERE / "fonts" / name)] + glob.glob(f"/usr/share/fonts/**/{name}", recursive=True)
        cands += glob.glob(f"C:/Windows/Fonts/{'arialbd' if k[0] else 'arial'}.ttf")
    path = next((c for c in cands if Path(c).exists()), None)
    _FONT_PATH[key] = path
    return path


def _truetype(path, size):
    try:
        return ImageFont.truetype(path, size)
    except (OSError, TypeError, ValueError):
        try:
            return ImageFont.truetype(path, max(1, int(round(size))))
        except Exception:                                          # noqa: BLE001
            return None


def font_px(size_px: float, bold: bool = False, italic: bool = False):
    """A font for drawing at `size_px` pixels."""
    path = _font_path(bold, italic)
    if path is None:
        return ImageFont.load_default()
    key = (round(size_px * 8) / 8, bool(bold), bool(italic))
    with _FONT_LOCK:
        if key not in _FONTS:
            if len(_FONTS) >= 48:                                          # every size is a separate FreeType face: keep the cache small
                _FONTS.clear()
            _FONTS[key] = _truetype(path, key[0]) or ImageFont.load_default()
        return _FONTS[key]


def text_w(text: str, size_pt: float, bold: bool = False, italic: bool = False) -> float:
    """Width in points (measured on a 10x font, so the answer does not depend on the drawing scale)."""
    f = font_px(size_pt * 10.0, bold, italic)
    try:
        return f.getlength(text) / 10.0
    except Exception:                                              # noqa: BLE001
        return len(text) * size_pt * 0.52


# ───────────────────────────────────────────────── helpers ─────────────────────────────────────────────────
def _num(v, default=0.0):
    try:
        return float(v)
    except (TypeError, ValueError):
        return default


def _hex(el):
    """colour of an a:solidFill-holding element -> 'RRGGBB' or None"""
    if el is None:
        return None
    c = el.find("a:solidFill/a:srgbClr", NS)
    if c is not None:
        return c.get("val")
    if el.find("a:solidFill/a:schemeClr", NS) is not None:      # theme colour: the engine never uses one; draw it dark
        return DEFAULT_COLOR
    return None


def _rgb(h):
    h = (h or DEFAULT_COLOR).lstrip("#")
    try:
        return tuple(int(h[i:i + 2], 16) for i in (0, 2, 4))
    except ValueError:
        return (35, 40, 45)


def _xfrm(el, path):
    x = el.find(path, NS)
    if x is None:
        return None
    off, ext = x.find("a:off", NS), x.find("a:ext", NS)
    if off is None or ext is None:
        return None
    return (int(off.get("x", 0)) / EMU, int(off.get("y", 0)) / EMU, int(ext.get("cx", 0)) / EMU, int(ext.get("cy", 0)) / EMU,
            x.get("flipH") == "1", x.get("flipV") == "1")


class _Run:
    __slots__ = ("text", "size", "bold", "italic", "color", "hl")

    def __init__(self, text, size, bold, italic, color, hl=None):
        self.text, self.size, self.bold, self.italic, self.color, self.hl = text, size, bold, italic, color, hl


class _Para:
    def __init__(self):
        self.items = []            # _Run | "BR"
        self.algn = "l"
        self.marL = 0.0
        self.indent = 0.0
        self.ln = None             # ("pts", v) | ("pct", v)
        self.bef = 0.0
        self.aft = 0.0
        self.bullet = None         # ("char", ch) | ("auto", type, startAt)
        self.end_size = DEFAULT_SIZE


# ───────────────────────────────────────────── text layout ─────────────────────────────────────────────
def _parse_paras(tx_body, dflt):
    """a:txBody -> [_Para].  `dflt`: the placeholder's lvl1 defaults {algn, ln, size, bold, italic, color}."""
    out = []
    for p in tx_body.findall("a:p", NS):
        para = _Para()
        ppr = p.find("a:pPr", NS)
        para.algn = {"ctr": "c", "r": "r", "l": "l", "just": "l"}.get((ppr.get("algn") if ppr is not None else None) or dflt.get("algn", "l"), "l")
        if ppr is not None:
            para.marL = _num(ppr.get("marL")) / EMU
            para.indent = _num(ppr.get("indent")) / EMU
            ln = ppr.find("a:lnSpc", NS)
            if ln is not None:
                if ln.find("a:spcPts", NS) is not None:
                    para.ln = ("pts", _num(ln.find("a:spcPts", NS).get("val")) / 100.0)
                elif ln.find("a:spcPct", NS) is not None:
                    para.ln = ("pct", _num(ln.find("a:spcPct", NS).get("val")) / 100000.0)
            for tag, attr in (("a:spcBef", "bef"), ("a:spcAft", "aft")):
                e = ppr.find(f"{tag}/a:spcPts", NS)
                if e is not None:
                    setattr(para, attr, _num(e.get("val")) / 100.0)
            if ppr.find("a:buAutoNum", NS) is not None:
                b = ppr.find("a:buAutoNum", NS)
                para.bullet = ("auto", b.get("type", "arabicPeriod"), int(_num(b.get("startAt"), 1)))
            elif ppr.find("a:buChar", NS) is not None:
                para.bullet = ("char", ppr.find("a:buChar", NS).get("char", "\u2022"))
        if para.ln is None and dflt.get("ln"):
            para.ln = dflt["ln"]
        for ch in p:
            tag = etree.QName(ch).localname
            if tag == "r":
                rpr = ch.find("a:rPr", NS)
                t = ch.find("a:t", NS)
                size = (_num(rpr.get("sz")) / 100.0) if rpr is not None and rpr.get("sz") else dflt.get("size", DEFAULT_SIZE)
                bold = (rpr.get("b") == "1") if rpr is not None and rpr.get("b") is not None else dflt.get("bold", False)
                ital = (rpr.get("i") == "1") if rpr is not None and rpr.get("i") is not None else dflt.get("italic", False)
                col = _hex(rpr) or dflt.get("color") or DEFAULT_COLOR
                hl = rpr.find("a:highlight/a:srgbClr", NS) if rpr is not None else None
                para.items.append(_Run((t.text or "") if t is not None else "", size, bold, ital, col, hl.get("val") if hl is not None else None))
            elif tag == "br":
                para.items.append("BR")
        epr = p.find("a:endParaRPr", NS)
        runs = [i for i in para.items if i != "BR"]
        para.end_size = ((_num(epr.get("sz")) / 100.0) if epr is not None and epr.get("sz") else
                         (runs[0].size if runs else dflt.get("size", DEFAULT_SIZE)))
        out.append(para)
    return out


def _atoms(para):
    """runs -> [('w', text, run) | ('s', ' ', run) | ('br',)]   words never break inside; NBSP is part of a word"""
    atoms = []
    for it in para.items:
        if it == "BR":
            atoms.append(("br",))
            continue
        buf, space = "", None
        for ch in it.text.replace("\t", " ").replace("\x0b", " "):
            if ch == " ":
                if buf:
                    atoms.append(("w", buf, it))
                    buf = ""
                atoms.append(("s", " ", it))
            else:
                buf += ch
        if buf:
            atoms.append(("w", buf, it))
    return atoms


def _split_long(word, run, avail):
    """a word wider than the line is cut between characters (as PowerPoint does)"""
    parts, cur = [], ""
    for ch in word:
        if cur and text_w(cur + ch, run.size, run.bold, run.italic) > avail:
            parts.append(cur)
            cur = ch
        else:
            cur += ch
    parts.append(cur)
    return parts


def _seg_width(atoms):
    """width of a run of atoms: consecutive atoms of one run are measured as ONE string, so the kerning across spaces
    ("A long") is included - that is how PowerPoint and LibreOffice measure a line"""
    total, buf, brun = 0.0, "", None
    for _kind, text, run in atoms:
        if brun is not None and run is brun:
            buf += text
        else:
            if buf:
                total += text_w(buf, brun.size, brun.bold, brun.italic)
            buf, brun = text, run
    if buf:
        total += text_w(buf, brun.size, brun.bold, brun.italic)
    return total


def _break_lines(para, avail_first, avail_rest, wrap=True):
    """-> [[(text, run)]]  (consecutive atoms of one run are merged)"""
    lines, cur, first = [], [], True

    def flush():
        nonlocal cur, first
        while cur and cur[-1][0] == "s":                         # trailing spaces hang outside the line
            cur.pop()
        lines.append(cur)
        cur, first = [], False

    for at in _atoms(para):
        if at[0] == "br":
            flush()
            continue
        kind, text, run = at
        avail = avail_first if first else avail_rest
        if kind == "s":
            if cur:
                cur.append(at)
            continue
        if wrap and cur and _seg_width(cur + [at]) > avail * FIT + 0.02:    # FIT: overflowing by less than 0.1 % still fits (as in LibreOffice)
            flush()
            avail = avail_rest
        w = text_w(text, run.size, run.bold, run.italic)
        if wrap and w > avail * FIT + 0.02 and not cur:
            parts = _split_long(text, run, max(avail, run.size))
            for k, part in enumerate(parts):
                cur.append(("w", part, run))
                if k < len(parts) - 1:
                    flush()
            continue
        cur.append(at)
    if cur or not lines:
        flush()
    out = []
    for ln in lines:
        segs = []
        for _k, text, run in ln:
            if segs and segs[-1][1] is run:
                segs[-1] = (segs[-1][0] + text, run)
            else:
                segs.append((text, run))
        out.append(segs)
    return out


def _auto_label(kind, n):
    if kind == "alphaUcPeriod":
        s, m = "", n
        while m > 0:
            m, r = divmod(m - 1, 26)
            s = chr(65 + r) + s
        return s + "."
    if kind == "alphaLcPeriod":
        return _auto_label("alphaUcPeriod", n).lower()
    if kind == "romanUcPeriod":
        vals = [(1000, "M"), (900, "CM"), (500, "D"), (400, "CD"), (100, "C"), (90, "XC"), (50, "L"), (40, "XL"), (10, "X"), (9, "IX"), (5, "V"), (4, "IV"), (1, "I")]
        s, m = "", n
        for v, t in vals:
            while m >= v:
                s += t
                m -= v
        return s + "."
    if kind == "arabicParenR":
        return f"{n})"
    return f"{n}."


def layout_text(paras, width, wrap=True):
    """-> (lines, total_height).  Each line: dict(y=top of the line, h, base=baseline y, x=left of the text, w, segs,
    label=(text, x, run) | None) with y / x relative to the top-left corner of the text area (points)."""
    lines, y = [], 0.0
    counters = {}                                                  # auto-number sequence: kind -> last number
    prev_auto = None
    for pi, para in enumerate(paras):
        has_b = para.bullet is not None
        text_left = para.marL if has_b else para.marL
        first_left = para.marL if has_b else para.marL + para.indent
        label_x = para.marL + para.indent
        avail_rest = max(5.0, width - para.marL)
        avail_first = max(5.0, width - first_left)
        if para.bef and pi > 0:
            y += para.bef
        brs = _break_lines(para, avail_first, avail_rest, wrap)
        label = None
        if has_b:
            if para.bullet[0] == "auto":
                kind, start = para.bullet[1], para.bullet[2]
                n = counters.get(kind, start - 1) + 1 if prev_auto == kind else start
                counters = {kind: n}
                prev_auto = kind
                label = _auto_label(kind, n)
            else:
                label = para.bullet[1]
                prev_auto = None
                counters = {}
        else:
            prev_auto = None
            counters = {}
        for li, segs in enumerate(brs):
            sizes = [r.size for _t, r in segs] or [para.end_size]
            size = max(sizes)
            single = SINGLE * size
            if para.ln is None:
                h = single
            elif para.ln[0] == "pts":
                h = para.ln[1]
            else:
                h = single * para.ln[1]
            base = h - DESC * size                                # the descender stays at the bottom of the line box
            left = (first_left if li == 0 else para.marL)
            wline = sum(text_w(t, r.size, r.bold, r.italic) for t, r in segs)
            area = width - left
            if para.algn == "c":
                x = left + (area - wline) / 2.0
            elif para.algn == "r":
                x = left + (area - wline)
            else:
                x = left
            lab = None
            if label is not None and li == 0:
                r0 = segs[0][1] if segs else _Run("", size, False, False, DEFAULT_COLOR)
                lab = (label, label_x, r0)
            lines.append(dict(y=y, h=h, base=y + base, x=x, w=wline, segs=segs, label=lab))
            y += h
        y += para.aft
    return lines, y


# ───────────────────────────────────────────── the deck ─────────────────────────────────────────────
class Deck:
    def __init__(self, path):
        self.path = str(path)
        self.prs = Presentation(self.path)
        self.W = self.prs.slide_width / EMU
        self.H = self.prs.slide_height / EMU
        self.unsupported = []
        self._blobs = {}               # key -> bytes
        self._dec = {}                 # key -> decoded RGBA
        self._scaled = {}              # (key, w, h, crop) -> RGBA
        self._lock = threading.Lock()

    # ----- information
    def count(self):
        return len(self.prs.slides)

    def title(self, i):
        sl = self.prs.slides[i]
        for sp in sl.shapes._spTree.findall("p:sp", NS):
            ph = sp.find("p:nvSpPr/p:nvPr/p:ph", NS)
            if ph is not None and ph.get("type") in ("title", "ctrTitle"):
                t = "".join(x.text or "" for x in sp.iter(f"{{{A}}}t")).strip()
                if t:
                    return t
        for sp in sl.shapes._spTree.findall("p:sp", NS):
            t = "".join(x.text or "" for x in sp.iter(f"{{{A}}}t")).strip()
            if t:
                return t
        return f"Slide {i + 1}"

    # ----- scene -> drawing operations (all in points)
    def _layout_ph(self, layout, ph):
        want_idx, want_type = ph.get("idx"), ph.get("type")
        best = None
        for sp in layout.shapes._spTree.findall("p:sp", NS):
            lp = sp.find("p:nvSpPr/p:nvPr/p:ph", NS)
            if lp is None:
                continue
            if want_idx is not None and lp.get("idx") == want_idx:
                return sp
            if lp.get("type") == want_type or (want_type in ("title", "ctrTitle") and lp.get("type") in ("title", "ctrTitle")):
                best = best if best is not None else sp
        return best

    def _blob(self, part, rid):
        try:
            rel = part.related_part(rid)
        except KeyError:
            return None
        blob = rel.blob
        key = hashlib.md5(blob).hexdigest()
        self._blobs.setdefault(key, blob)
        return key

    def _ops_for_tree(self, tree, part, layout, ops, is_layout):
        for el in tree:
            tag = etree.QName(el).localname
            if tag == "pic":
                xf = _xfrm(el, "p:spPr/a:xfrm")
                blip = el.find("p:blipFill/a:blip", NS)
                if xf is None or blip is None:
                    continue
                key = self._blob(part, blip.get(f"{{{RN}}}embed"))
                if key is None:
                    self.unsupported.append("picture without image data")
                    continue
                sr = el.find("p:blipFill/a:srcRect", NS)
                crop = tuple(_num(sr.get(k)) / 100000.0 for k in ("l", "t", "r", "b")) if sr is not None else (0, 0, 0, 0)
                ops.append(("pic", xf[0], xf[1], xf[2], xf[3], key, crop))
            elif tag == "sp":
                ph = el.find("p:nvSpPr/p:nvPr/p:ph", NS)
                if is_layout and ph is not None:
                    continue                                       # layout placeholders are only templates
                xf = _xfrm(el, "p:spPr/a:xfrm")
                lay_el, dflt, lay_body = None, {}, None
                if ph is not None and layout is not None:
                    lay_el = self._layout_ph(layout, ph)
                    if lay_el is not None:
                        if xf is None:
                            xf = _xfrm(lay_el, "p:spPr/a:xfrm")
                        lay_body = lay_el.find("p:txBody/a:bodyPr", NS)
                        lvl = lay_el.find("p:txBody/a:lstStyle/a:lvl1pPr", NS)
                        if lvl is not None:
                            d = lvl.find("a:defRPr", NS)
                            dflt = {"algn": lvl.get("algn", "l")}
                            ln = lvl.find("a:lnSpc/a:spcPts", NS)
                            if ln is not None:
                                dflt["ln"] = ("pts", _num(ln.get("val")) / 100.0)
                            if d is not None:
                                if d.get("sz"):
                                    dflt["size"] = _num(d.get("sz")) / 100.0
                                dflt["bold"] = d.get("b") == "1"
                                dflt["italic"] = d.get("i") == "1"
                                dflt["color"] = _hex(d)
                if xf is None:
                    continue
                x, y, w, h = xf[:4]
                prst = el.find("p:spPr/a:prstGeom", NS)
                if prst is not None and prst.get("prst", "rect") not in ("rect",):
                    self.unsupported.append(f"autoshape '{prst.get('prst')}' drawn as a rectangle")
                fill = _hex(el.find("p:spPr", NS))
                if fill:
                    ops.append(("rect", x, y, w, h, fill))
                ln = el.find("p:spPr/a:ln", NS)
                if ln is not None and ln.find("a:solidFill", NS) is not None and _num(ln.get("w")) > 0:
                    lw, lc = _num(ln.get("w")) / EMU, _hex(ln)
                    ops += [("rect", x, y, w, lw, lc), ("rect", x, y + h - lw, w, lw, lc), ("rect", x, y, lw, h, lc), ("rect", x + w - lw, y, lw, h, lc)]
                tx = el.find("p:txBody", NS)
                if tx is not None and "".join(t.text or "" for t in tx.iter(f"{{{A}}}t")).strip():
                    body = tx.find("a:bodyPr", NS)
                    props = {}
                    for src in (lay_body, body):                      # the shape's own values win over the layout's
                        if src is not None:
                            props.update({k: v for k, v in src.attrib.items()})
                    self._text_ops(ops, tx, x, y, w, h, props, dflt)
            elif tag == "cxnSp":
                xf = _xfrm(el, "p:spPr/a:xfrm")
                if xf is None:
                    continue
                x, y, w, h, fh, fv = xf
                ln = el.find("p:spPr/a:ln", NS)
                lw = (_num(ln.get("w"), 12700) / EMU) if ln is not None else 1.0
                lc = _hex(ln) or "000000"
                x1, x2 = (x + w, x) if fh else (x, x + w)
                y1, y2 = (y + h, y) if fv else (y, y + h)
                ops.append(("line", x1, y1, x2, y2, lw, lc))
            elif tag == "graphicFrame":
                xf = _xfrm(el, "p:xfrm")
                tbl = el.find("a:graphic/a:graphicData/a:tbl", NS)
                if xf is None or tbl is None:
                    self.unsupported.append("graphic frame that is not a table")
                    continue
                self._table_ops(ops, tbl, xf[0], xf[1])
            elif tag in ("nvGrpSpPr", "grpSpPr", "extLst"):
                continue
            elif tag == "grpSp":
                self.unsupported.append("grouped shapes")

    def _text_ops(self, ops, tx, x, y, w, h, props, dflt):
        l, t = _num(props.get("lIns"), 91440) / EMU, _num(props.get("tIns"), 45720) / EMU
        r, b = _num(props.get("rIns"), 91440) / EMU, _num(props.get("bIns"), 45720) / EMU
        wrap = props.get("wrap", "square") != "none"
        anchor = props.get("anchor", "t")
        paras = _parse_paras(tx, dflt)
        width = w - l - r
        lines, total = layout_text(paras, width, wrap)
        top = y + t
        if anchor == "ctr":
            top = y + t + ((h - t - b) - total) / 2.0
        elif anchor == "b":
            top = y + h - b - total
        self._line_ops(ops, lines, x + l, top)

    @staticmethod
    def _line_ops(ops, lines, x0, y0):
        for ln in lines:
            if ln["label"] is not None:
                text, lx, run = ln["label"]
                ops.append(("text", x0 + lx, y0 + ln["base"], text, run.size, run.bold, run.italic, run.color, True))
            x = x0 + ln["x"]
            for text, run in ln["segs"]:
                wseg = text_w(text, run.size, run.bold, run.italic)
                if text and run.hl:                                   # text highlight: a band behind the run (ascent .. descent of the font)
                    top = y0 + ln["base"] - 0.93 * run.size
                    ops.append(("rect", x, top, wseg, 1.17 * run.size, run.hl))
                if text:
                    ops.append(("text", x, y0 + ln["base"], text, run.size, run.bold, run.italic, run.color, False))
                x += wseg

    def _table_ops(self, ops, tbl, x0, y0):
        cols = [int(g.get("w", 0)) / EMU for g in tbl.findall("a:tblGrid/a:gridCol", NS)]
        rows = tbl.findall("a:tr", NS)
        # first pass: text layout -> the height each row really needs (rows grow with their text, like in PowerPoint)
        cells = []
        heights = []
        for tr in rows:
            row, need = [], int(tr.get("h", 0)) / EMU
            for ci, tc in enumerate(tr.findall("a:tc", NS)):
                pr = tc.find("a:tcPr", NS)
                mar = {k: _num(pr.get(k), d) / EMU if pr is not None else d / EMU for k, d in (("marL", 91440), ("marR", 91440), ("marT", 45720), ("marB", 45720))}
                tx = tc.find("a:txBody", NS)
                paras = _parse_paras(tx, {}) if tx is not None else []
                wcell = cols[ci] if ci < len(cols) else 100.0
                lines, total = layout_text(paras, wcell - mar["marL"] - mar["marR"], True)
                row.append((tc, pr, mar, lines, total))
                need = max(need, total + mar["marT"] + mar["marB"])
            cells.append(row)
            heights.append(need)
        ys = [y0]
        for hh in heights:
            ys.append(ys[-1] + hh)
        xs = [x0]
        for ww in cols:
            xs.append(xs[-1] + ww)
        for ri, row in enumerate(cells):                              # fills
            for ci, (tc, pr, mar, lines, total) in enumerate(row):
                fill = _hex(pr)
                if fill:
                    ops.append(("rect", xs[ci], ys[ri], cols[ci], heights[ri], fill))
        for ri, row in enumerate(cells):                              # borders (centred on the cell edges)
            for ci, (tc, pr, mar, lines, total) in enumerate(row):
                if pr is None:
                    continue
                for tag, edge in (("lnL", "L"), ("lnR", "R"), ("lnT", "T"), ("lnB", "B")):
                    ln = pr.find(f"a:{tag}", NS)
                    if ln is None or ln.find("a:solidFill", NS) is None:
                        continue
                    lw, lc = _num(ln.get("w"), 12700) / EMU, _hex(ln)
                    xa, xb, ya, yb = xs[ci], xs[ci] + cols[ci], ys[ri], ys[ri] + heights[ri]
                    if edge == "L":
                        ops.append(("rect", xa - lw / 2, ya - lw / 2, lw, yb - ya + lw, lc))
                    elif edge == "R":
                        ops.append(("rect", xb - lw / 2, ya - lw / 2, lw, yb - ya + lw, lc))
                    elif edge == "T":
                        ops.append(("rect", xa - lw / 2, ya - lw / 2, xb - xa + lw, lw, lc))
                    else:
                        ops.append(("rect", xa - lw / 2, yb - lw / 2, xb - xa + lw, lw, lc))
        for ri, row in enumerate(cells):                              # text
            for ci, (tc, pr, mar, lines, total) in enumerate(row):
                anchor = (pr.get("anchor") if pr is not None else None) or "t"
                top = ys[ri] + mar["marT"]
                room = heights[ri] - mar["marT"] - mar["marB"]
                if anchor == "ctr":
                    top += (room - total) / 2.0
                elif anchor == "b":
                    top += room - total
                self._line_ops(ops, lines, xs[ci] + mar["marL"], top)

    def ops(self, i):
        sl = self.prs.slides[i]
        layout = sl.slide_layout
        ops = []
        self._ops_for_tree(layout.shapes._spTree, layout.part, None, ops, True)
        self._ops_for_tree(sl.shapes._spTree, sl.part, layout, ops, False)
        return ops

    def text_lines(self, i):
        """[(x, baseline, text, size_pt, bold, is_bullet_label)] in points - what the self-test compares with LibreOffice"""
        return [(o[1], o[2], o[3], o[4], o[5], o[8]) for o in self.ops(i) if o[0] == "text"]

    # ----- pixels
    def _decoded(self, key):
        im = self._dec.get(key)
        if im is None:
            im = Image.open(io.BytesIO(self._blobs[key]))
            im.load()
            has_alpha = im.mode in ("RGBA", "LA", "PA") or (im.mode == "P" and "transparency" in im.info)
            im = im.convert("RGBA" if has_alpha else "RGB")              # opaque pictures stay RGB: a quarter less memory
            self._dec[key] = im
        return im

    def _scaled_pic(self, key, w, h, crop):
        k = (key, w, h, crop)
        im = self._scaled.get(k)
        if im is None:
            src = self._decoded(key)
            if any(crop):
                W0, H0 = src.size
                box = (round(crop[0] * W0), round(crop[1] * H0), W0 - round(crop[2] * W0), H0 - round(crop[3] * H0))
                src = src.crop(box)
            if src.mode == "RGBA":                                           # resample in premultiplied space: no dark fringe at transparent edges
                im = src.convert("RGBa").resize((max(1, w), max(1, h)), Image.LANCZOS, reducing_gap=2.0).convert("RGBA")
            else:
                im = src.resize((max(1, w), max(1, h)), Image.LANCZOS, reducing_gap=2.0)
            if len(self._scaled) > 8:                                        # slide pictures are cached on disk by the server; scaled copies are cheap to redo
                self._scaled.clear()
            self._scaled[k] = im
        return im

    def render(self, i, width_px=1280):
        """slide i -> PIL RGB image `width_px` wide"""
        width_px = int(max(120, min(width_px, 6000)))
        ss = 2 if width_px <= 1800 else 1
        k = width_px * ss / self.W
        cw, ch = round(self.W * k), round(self.H * k)
        canvas = Image.new("RGB", (cw, ch), (255, 255, 255))
        draw = ImageDraw.Draw(canvas)
        with self._lock:
            ops = self.ops(i)
        for op in ops:
            kind = op[0]
            if kind == "rect":
                _k, x, y, w, h, col = op
                draw.rectangle([round(x * k), round(y * k), max(round((x + w) * k) - 1, round(x * k)), max(round((y + h) * k) - 1, round(y * k))],
                               fill=_rgb(col))
            elif kind == "line":
                _k, x1, y1, x2, y2, lw, col = op
                if abs(x1 - x2) < 1e-6:                                    # vertical: a filled bar (crisp, exact width)
                    ya, yb = sorted((y1, y2))
                    x_a = round((x1 - lw / 2) * k)
                    draw.rectangle([x_a, round(ya * k), max(round((x1 + lw / 2) * k) - 1, x_a), round(yb * k)], fill=_rgb(col))
                elif abs(y1 - y2) < 1e-6:                                  # horizontal
                    xa, xb = sorted((x1, x2))
                    y_a = round((y1 - lw / 2) * k)
                    draw.rectangle([round(xa * k), y_a, round(xb * k), max(round((y1 + lw / 2) * k) - 1, y_a)], fill=_rgb(col))
                else:
                    draw.line([x1 * k, y1 * k, x2 * k, y2 * k], fill=_rgb(col), width=max(1, round(lw * k)))
            elif kind == "pic":
                _k, x, y, w, h, key, crop = op
                px, py, pw, ph = round(x * k), round(y * k), round(w * k), round(h * k)
                if pw < 1 or ph < 1:
                    continue
                with self._lock:
                    if key not in self._dec:
                        self._decoded(key)
                try:
                    pic = self._scaled_pic(key, pw, ph, crop)
                except Exception:                                              # noqa: BLE001 - one bad picture never stops the preview
                    self.unsupported.append("a picture could not be decoded")
                    continue
                canvas.paste(pic, (px, py), pic if pic.mode == "RGBA" else None)
            elif kind == "text":
                _k, x, base, text, size, bold, ital, col, _lab = op
                f = font_px(size * k, bold, ital)
                txt = text.replace("\u00a0", " ")
                try:
                    draw.text((x * k, base * k), txt, font=f, fill=_rgb(col), anchor="ls")
                except (ValueError, OSError, TypeError):                     # a font without anchor support: place by the ascent
                    draw.text((x * k, base * k - size * k * 0.9), txt, font=f, fill=_rgb(col))
        if ss > 1:
            canvas = canvas.resize((width_px, round(self.H * width_px / self.W)), Image.LANCZOS)
        return canvas


# ───────────────────────────────────────────── caching + PDF ─────────────────────────────────────────────
_DECKS: dict = {}
_DECKS_LOCK = threading.Lock()


def get_deck(path) -> Deck:
    """a parsed deck, reused while the file is unchanged (the web page asks for the slides one by one)"""
    p = Path(path)
    key = (str(p), p.stat().st_mtime_ns, p.stat().st_size)
    with _DECKS_LOCK:
        d = _DECKS.get(key)
        if d is None:
            if len(_DECKS) >= 2:
                _DECKS.pop(next(iter(_DECKS)))
            d = _DECKS[key] = Deck(p)
        return d


def images_to_pdf(deck: Deck, out_pdf, width_px=1920, quality=90):
    """A PDF made of the slide pictures (page = 960 x 540 pt).  Text cannot be selected in it - it is the fallback for
    computers without LibreOffice.  Pages are kept as compressed JPEGs until the file is written (a long deck stays small in memory)."""
    import io as _io
    n = deck.count()
    if n == 0:
        raise ValueError("the presentation has no slides")
    pages = []
    for i in range(n):
        buf = _io.BytesIO()
        deck.render(i, width_px).convert("RGB").save(buf, "JPEG", quality=95, subsampling=0)
        buf.seek(0)
        pages.append(Image.open(buf))                               # opened, not decoded: decoded one by one while the PDF is written
    res = pages[0].width / (deck.W / 72.0)
    pages[0].save(str(out_pdf), "PDF", save_all=True, append_images=pages[1:], resolution=res, quality=quality)
    return str(out_pdf)


if __name__ == "__main__":
    if len(sys.argv) < 3:
        sys.exit(__doc__)
    d = Deck(sys.argv[1])
    out = Path(sys.argv[2])
    out.mkdir(parents=True, exist_ok=True)
    wpx = int(sys.argv[3]) if len(sys.argv) > 3 else 1280
    for n in range(d.count()):
        d.render(n, wpx).save(out / f"s{n + 1:02d}.png")
        print(f"slide {n + 1}: {d.title(n)}")
    if d.unsupported:
        print("not drawn:", sorted(set(d.unsupported)))
