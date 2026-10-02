#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
mw_report.py  -  Mech Well FEA report builder  (python-pptx)

Re-creates the "Static Structural Analysis" PPT report template (sample.pdf) as an
editable .pptx.  Every coordinate below was measured from the template PDF
(page = 960 x 540 pt = 13.333 x 7.5 in, 16:9).

Slide recipe (the "procedure"):
    1  Cover ................ FEA REPORT label, title, Report No., Date, Client, Mech Well logo
    2  1. Material Properties  table + units + sigma_all = Syt / FOS
    3  2. <Object> Geometry ... picture
    4  3. <Object> Mesh ....... picture
    per load case:
        Case N: <name> (Self Weight + Pressure) ... BC picture | Boundary Conditions A-D | Note
        Results ............... Total deformation | Von-Mises stress (2 pictures + captions)
        Observation & Summary . 3 bullets (+ corrective-action line when stress > allowable)
    last  Final Summary ...... table: Load Case | Max deformation | Max Von-Mises | Allowable

Usage:
    python3 mw_report.py config.json            # -> pptx path given in config["output"]
    (or)  from mw_report import build_report; build_report(cfg_dict)
"""
from __future__ import annotations

import glob
import hashlib
import io
import json
import os
import re
import sys
import tempfile
from pathlib import Path

from PIL import Image, ImageChops, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Emu, Pt

HERE = Path(__file__).resolve().parent
ASSETS = HERE / "assets"

# ───────────────────────── design tokens (measured from the template) ─────────────────────────
W, H = 960.0, 540.0
FONT = "Arial"
NAVY, BLUE, GREY, BODY = "1B304E", "2A6FA6", "5F6973", "23282D"
TEAL_DK, TEAL = "023C51", "03506C"
ROW_LT, BAND1, BAND2, TXT_GREY, WHITE = "F2F5F8", "CBD0D4", "E7E9EB", "595959", "FFFFFF"

# picture slots (x0, y0, x1, y1) in points
GEO_BOX = (121.0, 84.5, 839.0, 510.0)      # geometry / mesh  (centred)
BC_IMG_BOX = (24.0, 88.0, 564.0, 483.0)    # boundary-condition picture (left aligned, template: x=17..24, h=397)
RES_L_BOX = (74.0, 100.0, 466.0, 470.0)    # results, left column
RES_R_BOX = (494.0, 100.0, 886.0, 470.0)   # results, right column (kept clear of corner triangle)

LINE_BC = 17.7       # exact line pitch used by the template for 14 pt lists
LINE_NOTE = 16.8     # 14 pt notes
LINE_OBS = 21.5      # 17 pt observation bullets

MAX_SIDE = 2600      # px - bigger screenshots are scaled down (keeps the PowerPoint small; 2600 px on 880 pt = 213 dpi)
STK_X0, STK_X1 = 54.0, 906.0           # content area for stacked / one-per-slide layouts
STK_TOP, STK_BOT = 84.0, 524.0
HEAD_STACK = 22.0                      # heading height in the stacked layout
CAP_GAP_STACK = 6.0
SUMMARY_ROWS = 8                       # rows per summary slide (more -> a continuation slide)


def E(pt: float) -> Emu:
    return Emu(int(round(pt * 12700)))


# ───────────────────────────────── text measurement ─────────────────────────────────
_FONT_CACHE: dict = {}


def _font(size: float, bold: bool):
    key = (round(size * 10), bold)
    if key in _FONT_CACHE:
        return _FONT_CACHE[key]
    pat = "LiberationSans-Bold.ttf" if bold else "LiberationSans-Regular.ttf"
    # bundled copy first (metric-compatible with Arial, SIL OFL) -> identical wrapping on every computer
    paths = [str(HERE / "fonts" / pat)] + glob.glob(f"/usr/share/fonts/**/{pat}", recursive=True)
    paths = [q for q in paths if Path(q).exists()]
    f = ImageFont.truetype(paths[0], int(round(size * 10))) if paths else None
    _FONT_CACHE[key] = f
    return f


def text_width(text: str, size: float, bold: bool = False) -> float:
    f = _font(size, bold)
    if f is None:                       # fallback: average Arial glyph ~0.52 em
        return len(text) * size * (0.56 if bold else 0.52)
    return f.getlength(text) / 10.0


def count_lines(text: str, width: float, size: float, bold: bool = False) -> int:
    lines, cur = 1, ""
    for word in text.split():
        trial = f"{cur} {word}".strip()
        if cur and text_width(trial, size, bold) > width:
            lines += 1
            cur = word
        else:
            cur = trial
    return lines


_UNITS = r"(?:mm|cm|m|MPa|GPa|kPa|Pa|N|kN|kg|t|mmWC|Hz|s|rpm|\u00b0C|%)"


def glue_units(text: str) -> str:
    """Keep a number and its unit together when a line wraps (non-breaking space)."""
    return re.sub(rf"(?<=\d) (?={_UNITS}(?![A-Za-z]))", "\u00a0", text)



def balanced_lines(text: str, max_w: float, size: float, bold: bool = False):
    """Single line if it fits, else the most even 2-line split (3+ lines: greedy)."""
    if text_width(text, size, bold) <= max_w:
        return [text]
    words = text.split()
    best = None
    for i in range(1, len(words)):
        a, b = " ".join(words[:i]), " ".join(words[i:])
        wa, wb = text_width(a, size, bold), text_width(b, size, bold)
        if wa <= max_w and wb <= max_w and (best is None or max(wa, wb) < best[0]):
            best = (max(wa, wb), [a, b])
    if best:
        return best[1]
    lines, cur = [], ""
    for w in words:
        t = f"{cur} {w}".strip()
        if cur and text_width(t, size, bold) > max_w:
            lines.append(cur)
            cur = w
        else:
            cur = t
    return lines + [cur]


# ───────────────────────────────── low-level helpers ─────────────────────────────────
def R(text, size, bold=False, color=BODY, italic=False):
    return dict(text=text, size=size, bold=bold, color=color, italic=italic)


def P(runs, **kw):
    if isinstance(runs, dict):
        runs = [runs]
    return dict(runs=runs, **kw)


_ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
_ANCHOR = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}
_PPR_AFTER_BU = ("a:tabLst", "a:defRPr", "a:extLst")


# A placeholder the engineer still has to fill in is marked with a yellow highlight, so it cannot be overlooked in PowerPoint.
_PH_RE = re.compile(r"(\[[^\[\]]*(?:please confirm|not entered)[^\[\]]*\])", re.I)
HIGHLIGHT = "FFFF00"


def highlight_run(r):
    """<a:highlight> sits after the fill and before the font in a:rPr (ECMA-376 21.1.2.3.9)"""
    hl = parse_xml(f'<a:highlight {nsdecls("a")}><a:srgbClr val="{HIGHLIGHT}"/></a:highlight>')
    r._r.get_or_add_rPr().insert_element_before(hl, "a:uLnTx", "a:uLn", "a:uFillTx", "a:uFill", "a:latin", "a:ea", "a:cs", "a:sym",
                                                 "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst")


def style_run(r, spec):
    f = r.font
    f.name = FONT
    f.size = Pt(spec["size"])
    f.bold = bool(spec.get("bold"))
    f.italic = bool(spec.get("italic"))
    f.color.rgb = RGBColor.from_string(spec["color"])
    r._r.get_or_add_rPr().set("lang", "en-US")


def fill_paragraph(para, spec):
    if "align" in spec:
        para.alignment = _ALIGN[spec["align"]]
    if spec.get("line_pts"):
        para.line_spacing = Pt(spec["line_pts"])
    if spec.get("after") is not None:
        para.space_after = Pt(spec["after"])
    if spec.get("before") is not None:
        para.space_before = Pt(spec["before"])
    pPr = para._p.get_or_add_pPr()
    if "marL" in spec:
        pPr.set("marL", str(int(E(spec["marL"]))))
    if "indent" in spec:
        pPr.set("indent", str(int(E(spec["indent"]))))
    bullet = spec.get("bullet")
    if bullet:
        buFont = parse_xml(f'<a:buFont {nsdecls("a")} typeface="{FONT}"/>')
        pPr.insert_element_before(buFont, "a:buNone", "a:buAutoNum", "a:buChar", "a:buBlip", *_PPR_AFTER_BU)
        if bullet == "char":
            bu = parse_xml(f'<a:buChar {nsdecls("a")} char="&#8226;"/>')
        else:  # 'alphaUcPeriod' | 'arabicPeriod'
            bu = parse_xml(f'<a:buAutoNum {nsdecls("a")} type="{bullet}"/>')
        pPr.insert_element_before(bu, "a:buBlip", *_PPR_AFTER_BU)
    for item in spec["runs"]:
        if item == "BR":
            para.add_line_break()
            continue
        txt = item["text"]
        for part in (_PH_RE.split(txt) if _PH_RE.search(txt) else [txt]):      # "[... please confirm]" gets its own, highlighted run
            if part == "" and txt != "":
                continue
            r = para.add_run()
            r.text = part
            style_run(r, item)
            if _PH_RE.fullmatch(part):
                highlight_run(r)


def add_text(slide, x, y, w, h, paras, anchor="t", wrap=True, name=None):
    tb = slide.shapes.add_textbox(E(x), E(y), E(w), E(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = _ANCHOR[anchor]
    for i, spec in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        fill_paragraph(para, spec)
    return tb


# ───────────────────────────────────── images ─────────────────────────────────────
_PH_DIR = Path(tempfile.gettempdir()) / "fea_report_maker_placeholders"


def placeholder_picture(title: str, hint: str = "Right-click this box > Change Picture to insert your picture here.",
                        size=(1000, 760)) -> dict:
    """An empty labelled frame (a real picture, so it can be replaced in PowerPoint with Change Picture and keeps its place).
    -> image spec (path + autocrop off)."""
    from PIL import ImageDraw
    key = hashlib.md5(f"{title}|{hint}|{size}".encode("utf-8")).hexdigest()[:12]
    path = _PH_DIR / f"ph_{key}.png"
    if not path.exists():
        _PH_DIR.mkdir(parents=True, exist_ok=True)
        w, h = size
        im = Image.new("RGB", (w, h), (238, 242, 246))
        d = ImageDraw.Draw(im)
        dash, gap, bw = 26, 16, 4
        col = (154, 168, 184)
        for x in range(0, w, dash + gap):                       # dashed border
            d.rectangle([x, 0, min(x + dash, w), bw], fill=col)
            d.rectangle([x, h - bw - 1, min(x + dash, w), h - 1], fill=col)
        for y in range(0, h, dash + gap):
            d.rectangle([0, y, bw, min(y + dash, h)], fill=col)
            d.rectangle([w - bw - 1, y, w - 1, min(y + dash, h)], fill=col)
        blocks = []                                               # (lines, font, colour, pitch): title, then the hint - stacked, never overlapping
        for txt, frac, bold, rgb in ((title, 0.07 if w >= 1100 else 0.062, True, (95, 105, 115)), (hint, 0.034, False, (125, 135, 146))):
            f = _font(h * frac / 10.0, bold)
            if f is None:
                continue
            lines, cur = [], ""
            for wd in txt.split():                                # wrap to 82 % of the width
                t = f"{cur} {wd}".strip()
                if cur and f.getlength(t) > 0.82 * w:
                    lines.append(cur)
                    cur = wd
                else:
                    cur = t
            lines.append(cur)
            blocks.append((lines, f, rgb, h * frac * 1.28))
        gap = h * 0.035
        total = sum(len(ln) * pitch for ln, _f, _c, pitch in blocks) + gap * max(0, len(blocks) - 1)
        y = (h - total) / 2
        for lines, f, rgb, pitch in blocks:
            for ln in lines:
                d.text(((w - f.getlength(ln)) / 2, y), ln, font=f, fill=rgb)
                y += pitch
            y += gap
        im.save(path, "PNG")
    return {"path": str(path), "autocrop": False}


_IMG_CACHE: dict = {}                       # (path, mtime, size, autocrop, pad) -> (png/jpeg bytes, (w, h), k, is_png)
_IMG_CACHE_BYTES = [0]
_IMG_CACHE_LIMIT = 160 * 1024 * 1024


def _cache_key(spec, autocrop, pad_frac):
    ac, p = autocrop, spec
    if isinstance(spec, dict):
        ac, p = spec.get("autocrop", autocrop), spec["path"]
    try:
        st = os.stat(p)
        return (str(p), st.st_mtime_ns, st.st_size, bool(ac), pad_frac)
    except OSError:
        return None


def _prepare_image(spec, autocrop=True, pad_frac=0.012):
    """-> (PIL image as it will be placed, k) : white margins trimmed, huge pictures scaled down by k (k = 1 when untouched)."""
    if isinstance(spec, dict):
        autocrop = spec.get("autocrop", autocrop)
        spec = spec["path"]
    im = Image.open(spec)
    im.load()
    fmt_png = True
    if im.mode in ("RGBA", "LA", "P"):
        im = im.convert("RGBA")
    else:
        im = im.convert("RGB")
        fmt_png = Path(spec).suffix.lower() == ".png"
    if autocrop:
        flat = im if im.mode == "RGB" else Image.alpha_composite(Image.new("RGBA", im.size, (255, 255, 255, 255)), im).convert("RGB")
        bg = Image.new("RGB", flat.size, flat.getpixel((0, 0)))
        diff = ImageChops.difference(flat, bg).convert("L").point(lambda v: 255 if v > 14 else 0)
        bbox = diff.getbbox()
        if bbox:
            pad = max(4, int(pad_frac * max(im.size)))
            l, t, r, b = bbox
            box = (max(0, l - pad), max(0, t - pad), min(im.width, r + pad), min(im.height, b + pad))
            if (box[2] - box[0]) * (box[3] - box[1]) < 0.97 * im.width * im.height:
                im = im.crop(box)
    k = 1.0
    if max(im.size) > MAX_SIDE:                                   # huge screenshot -> keep the file small
        k = MAX_SIDE / max(im.size)
        im = im.resize((max(1, round(im.width * k)), max(1, round(im.height * k))), Image.LANCZOS)
    im.info["_fmt_png"] = fmt_png
    return im, k


def load_image(spec, autocrop=True, pad_frac=0.012):
    """spec: path | {"path":..., "autocrop": bool}.  Returns (BytesIO, (w_px, h_px)).
    Auto-crop trims uniform (white) margins so every picture can be fitted at max size."""
    key = _cache_key(spec, autocrop, pad_frac)
    if key is not None and key in _IMG_CACHE:
        data, size, _k, _png = _IMG_CACHE[key]
        return io.BytesIO(data), size
    im, _k = _prepare_image(spec, autocrop, pad_frac)
    fmt_png = im.info.pop("_fmt_png", True)
    buf = io.BytesIO()
    if fmt_png or im.mode == "RGBA":
        im.save(buf, "PNG", optimize=True)
    else:
        im.save(buf, "JPEG", quality=92)
    data = buf.getvalue()
    if key is not None and len(data) < _IMG_CACHE_LIMIT // 4:
        if _IMG_CACHE_BYTES[0] + len(data) > _IMG_CACHE_LIMIT:               # simple reset instead of LRU bookkeeping
            _IMG_CACHE.clear()
            _IMG_CACHE_BYTES[0] = 0
        _IMG_CACHE[key] = (data, im.size, _k, fmt_png)
        _IMG_CACHE_BYTES[0] += len(data)
    buf.seek(0)
    return buf, im.size


def _tri_edges():
    """Left edge (pt) of the blue corner triangle for every pt-row, measured from its alpha mask."""
    im = Image.open(ASSETS / "triangle_corner.png").convert("RGBA")
    px = im.getchannel("A").load()
    edges = []
    for r in range(im.height):
        xs = [c for c in range(im.width) if px[c, r] > 128]
        edges.append(828.0 + min(xs) if xs else None)
    return edges


_TRI = _tri_edges()


def tri_left(y: float) -> float:
    r = int(y - 141.72)
    if r < 0:
        return 1e9
    e = _TRI[min(r, len(_TRI) - 1)]
    return 1e9 if e is None else e


def fit(size_px, box, ax="c", ay="t", margin=8.0, max_scale=None):
    """Largest rect with the picture's aspect ratio that fits `box`, shrunk further if its
    bottom-right corner would touch the corner triangle.  -> (x, y, w, h, scale_pt_per_px)"""
    x0, y0, x1, y1 = box
    bw, bh = x1 - x0, y1 - y0
    s = min(bw / size_px[0], bh / size_px[1], max_scale or 1e9)
    for _ in range(200):
        w, h = size_px[0] * s, size_px[1] * s
        x = x0 + {"l": 0, "c": (bw - w) / 2, "r": bw - w}[ax]
        y = y0 + {"t": 0, "c": (bh - h) / 2, "b": bh - h}[ay]
        if x + w <= tri_left(y + h) - margin:
            break
        s *= 0.99
    return x, y, w, h, s


def place_picture(slide, buf, rect, name=None, alt=None):
    x, y, w, h = rect[:4]
    pic = slide.shapes.add_picture(buf, E(x), E(y), E(w), E(h))
    if name:
        pic.name = name
    if alt:
        pic._element.nvPicPr.cNvPr.set("descr", alt)
    return pic


def add_picture(slide, spec, box, autocrop, ax="c", ay="t", name=None, alt=None):
    buf, size = load_image(spec, autocrop)
    rect = fit(size, box, ax, ay)
    return place_picture(slide, buf, rect, name, alt), rect[:4]


# ──────────────────────────────── layouts / branding ────────────────────────────────
def _next_id(spTree) -> int:
    ids = [int(i) for i in spTree.xpath(".//p:cNvPr/@id")] or [1]
    return max(ids) + 1


def layout_add_picture(layout, img_path, x, y, w, h, name):
    _, rId = layout.part.get_or_add_image_part(str(img_path))
    spTree = layout.shapes._spTree
    xml = f"""<p:pic {nsdecls("p", "a", "r")}>
      <p:nvPicPr><p:cNvPr id="{_next_id(spTree)}" name="{name}"/>
        <p:cNvPicPr><a:picLocks noChangeAspect="1"/></p:cNvPicPr><p:nvPr userDrawn="1"/></p:nvPicPr>
      <p:blipFill><a:blip r:embed="{rId}"/><a:stretch><a:fillRect/></a:stretch></p:blipFill>
      <p:spPr><a:xfrm><a:off x="{int(E(x))}" y="{int(E(y))}"/><a:ext cx="{int(E(w))}" cy="{int(E(h))}"/></a:xfrm>
        <a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr></p:pic>"""
    spTree.append(parse_xml(xml))


def layout_add_rect(layout, x, y, w, h, color, name):
    spTree = layout.shapes._spTree
    xml = f"""<p:sp {nsdecls("p", "a")}>
      <p:nvSpPr><p:cNvPr id="{_next_id(spTree)}" name="{name}"/><p:cNvSpPr/><p:nvPr userDrawn="1"/></p:nvSpPr>
      <p:spPr><a:xfrm><a:off x="{int(E(x))}" y="{int(E(y))}"/><a:ext cx="{int(E(w))}" cy="{int(E(h))}"/></a:xfrm>
        <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
        <a:solidFill><a:srgbClr val="{color}"/></a:solidFill><a:ln><a:noFill/></a:ln></p:spPr>
      <p:txBody><a:bodyPr rtlCol="0" anchor="ctr"/><a:lstStyle/><a:p><a:endParaRPr lang="en-US"/></a:p></p:txBody></p:sp>"""
    spTree.append(parse_xml(xml))


def _style_title_placeholder(layout, x, y, w, h, size, line_pts=None):
    ph = [p for p in layout.placeholders if p.placeholder_format.type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)][0]
    ph.left, ph.top, ph.width, ph.height = E(x), E(y), E(w), E(h)
    txBody = ph._element.txBody
    bodyPr = txBody.find(qn("a:bodyPr"))
    for ch in list(bodyPr):
        bodyPr.remove(ch)
    for k, v in dict(lIns="0", tIns="0", rIns="0", bIns="0", anchor="t", wrap="square").items():
        bodyPr.set(k, v)
    bodyPr.append(parse_xml(f'<a:noAutofit {nsdecls("a")}/>'))
    lst = txBody.find(qn("a:lstStyle"))
    for ch in list(lst):
        lst.remove(ch)
    ln = f'<a:lnSpc><a:spcPts val="{int(line_pts * 100)}"/></a:lnSpc>' if line_pts else ""
    lst.append(parse_xml(
        f'<a:lvl1pPr {nsdecls("a")} algn="l">{ln}<a:defRPr sz="{int(size * 100)}" b="1">'
        f'<a:solidFill><a:srgbClr val="{NAVY}"/></a:solidFill><a:latin typeface="{FONT}"/><a:cs typeface="{FONT}"/>'
        f'</a:defRPr></a:lvl1pPr>'))


def build_template() -> Presentation:
    prs = Presentation()
    prs.slide_width, prs.slide_height = E(W), E(H)

    master = prs.slide_master
    for shp in master.shapes:                       # default template is 4:3 -> stretch to 16:9
        shp.left, shp.width = int(shp.left * 4 / 3), int(shp.width * 4 / 3)

    # keep only the two layouts we need
    for lay in list(prs.slide_layouts):
        if lay.name not in ("Title Slide", "Title Only"):
            prs.slide_layouts.remove(lay)
    cover = next(l for l in prs.slide_layouts if l.name == "Title Slide")
    content = next(l for l in prs.slide_layouts if l.name == "Title Only")
    cover.name, content.name = "MW Cover", "MW Content"

    drop = (PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.SUBTITLE)
    for lay in (cover, content):
        for ph in list(lay.placeholders):
            if ph.placeholder_format.type in drop:
                ph._element.getparent().remove(ph._element)

    # theme: Arial + brand colours (so anything the user adds later matches)
    theme = master.part.part_related_by(RT.THEME)
    xml = theme.blob.decode("utf-8")
    xml = re.sub(r'(<a:(?:major|minor)Font>\s*<a:latin typeface=")[^"]*(")', rf"\g<1>{FONT}\2", xml)
    for tag, col in (("dk2", NAVY), ("accent1", BLUE), ("accent2", NAVY), ("accent3", TEAL), ("accent4", GREY)):
        xml = re.sub(rf'(<a:{tag}>\s*<a:srgbClr val=")[0-9A-Fa-f]{{6}}(")', rf"\g<1>{col}\2", xml)
    theme._blob = xml.encode("utf-8")

    # ── content layout: title, corner triangle, small logo, title rule (+ soft shadow)
    _style_title_placeholder(content, 53.3, 31.1, 800.0, 36.0, 25)
    layout_add_picture(content, ASSETS / "triangle_corner.png", 828.0, 141.72, 132.0, 399.0, "Corner triangle")
    layout_add_picture(content, ASSETS / "logo_small_topright.png", 904.56, 0.96, 55.44, 55.56, "Mech Well logo")
    layout_add_picture(content, ASSETS / "title_line_shadow.png", 47.04, 70.44, 863.16, 8.88, "Title rule shadow")
    layout_add_rect(content, 50.4, 72.0, 856.8, 2.5, BLUE, "Title rule")

    # ── cover layout: triangle, small logo, top bar (+shadow), big logo
    _style_title_placeholder(cover, 60.5, 97.8, 650.0, 80.0, 30, line_pts=36.0)
    layout_add_picture(cover, ASSETS / "triangle_corner.png", 828.0, 141.72, 132.0, 399.0, "Corner triangle")
    layout_add_picture(cover, ASSETS / "logo_small_topright.png", 904.56, 0.96, 55.44, 55.56, "Mech Well logo (small)")
    layout_add_picture(cover, ASSETS / "topbar_shadow_cover.png", -3.24, -1.56, 966.12, 13.44, "Top bar shadow")
    layout_add_rect(cover, 0.0, 0.0, 960.0, 7.1, BLUE, "Top bar")
    layout_add_picture(cover, ASSETS / "logo_mechwell_large.png", 535.44, 313.2, 325.92, 86.4, "Mech Well logo")
    return prs


def _layout(prs, name):
    return next(l for l in prs.slide_layouts if l.name == name)


def content_slide(prs, title, sub=None):
    s = prs.slides.add_slide(_layout(prs, "MW Content"))
    p = s.shapes.title.text_frame.paragraphs[0]
    for k in (1.0, 0.88, 0.8, 0.72, 0.64):                     # long titles shrink (25 pt -> 16 pt) instead of running off the slide
        w = text_width(title, 25 * k, True) + ((text_width(" ", 25 * k, True) + text_width(sub, 16 * k, True)) if sub else 0.0)
        if w <= 790.0:
            break
    while w > 790.0 and len(title) > 12:                        # still too long: shorten with an ellipsis
        title = title[:-2].rstrip() + "\u2026"
        w = text_width(title, 25 * k, True) + ((text_width(" ", 25 * k, True) + text_width(sub, 16 * k, True)) if sub else 0.0)
    runs = [R(title, 25 * k, True, NAVY)]
    if sub:
        runs += [R(" ", 25 * k, True, NAVY), R(sub, 16 * k, True, NAVY)]
    fill_paragraph(p, P(runs))
    return s


# ─────────────────────────────────────── tables ───────────────────────────────────────
def _set_borders(cell, spec):
    """spec: {'L'|'R'|'T'|'B': (width_pt, 'RRGGBB')}  -> inserted BEFORE the fill (schema order)."""
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for el in tcPr.findall(qn(tag)):
            tcPr.remove(el)
    new = []
    for tag, key in (("a:lnL", "L"), ("a:lnR", "R"), ("a:lnT", "T"), ("a:lnB", "B")):
        if key in spec:
            w, col = spec[key]
            new.append(parse_xml(
                f'<{tag} {nsdecls("a")} w="{int(E(w))}" cap="flat" cmpd="sng" algn="ctr">'
                f'<a:solidFill><a:srgbClr val="{col}"/></a:solidFill><a:prstDash val="solid"/>'
                f'<a:round/><a:headEnd type="none" w="med" len="med"/><a:tailEnd type="none" w="med" len="med"/></{tag}>'))
    for i, el in enumerate(new):
        tcPr.insert(i, el)


def add_table(slide, x, y, col_w, row_h, rows, header_fill, band_fills, hdr_size, body_size, body_color, name, top_margin=3.6):
    nrows, ncols = len(rows), len(col_w)
    gf = slide.shapes.add_table(nrows, ncols, E(x), E(y), E(sum(col_w)), E(sum(row_h)))
    gf.name = name
    tbl = gf.table
    for j, w in enumerate(col_w):
        tbl.columns[j].width = E(w)
    for i, h in enumerate(row_h):
        tbl.rows[i].height = E(h)
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "1")
    tblPr.set("bandRow", "0")
    tblPr.find(qn("a:tableStyleId")).text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"   # No Style, No Grid
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            head = i == 0
            cell.margin_left = cell.margin_right = E(7.2)
            cell.margin_top = E(top_margin)
            cell.margin_bottom = E(3.6)
            cell.vertical_anchor = MSO_ANCHOR.TOP
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(header_fill if head else band_fills[(i - 1) % len(band_fills)])
            tf = cell.text_frame
            tf.word_wrap = True
            for k, line in enumerate(str(val).split("\n")):
                para = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                spec = R(line, hdr_size if head else body_size, head, WHITE if head else body_color)
                fill_paragraph(para, P([spec], align="c"))
            b = {"L": (1, WHITE), "R": (1, WHITE), "T": (1, WHITE), "B": (1, WHITE)}
            if head:
                b["B"] = (3, WHITE)
            if i == 1:
                b["T"] = (3, WHITE)
            _set_borders(cell, b)
    return gf


# ────────────────────────────────── slide builders ──────────────────────────────────
def build_cover(prs, c):
    s = prs.slides.add_slide(_layout(prs, "MW Cover"))
    add_text(s, 60.5, 62.3, 300, 18, [P([R(c.get("label", "FEA REPORT"), 12, True, BLUE)])], wrap=False, name="Report label")
    title = c["title"].replace(" & ", " &\u00a0")      # keep '& Frame' together when wrapping
    fill_paragraph(s.shapes.title.text_frame.paragraphs[0], P([R(title, 30, True, NAVY)], line_pts=36.0))
    n = sum(count_lines(part, 650.0, 30, True) for part in title.split("\n"))
    dy = (min(n, 3) - 2) * 36.0                        # template has a 2-line title; shift details for 1 or 3 lines
    add_text(s, 60.5, 183.5 + dy, 560, 20, [P([R(f"Report No.: {c['report_no']}", 15, False, GREY)])], wrap=False, name="Report No")
    add_text(s, 60.5, 216.1 + dy, 560, 20, [P([R(f"Date: {c['date']}", 15, False, GREY)])], wrap=False, name="Date")
    add_text(s, 60.5, 255.5 + dy, 560, 18,
             [P([R("Client: ", 12, True, GREY), R(c["client"], 9, True, TXT_GREY)])], wrap=False, name="Client")
    # optional lines (only when filled in): revision / prepared / checked / drawing reference
    y = 255.5 + dy + 30.0
    sign = [f"{lab}: {c[key]}" for lab, key in (("Rev.", "revision"), ("Prepared by", "prepared_by"), ("Checked by", "checked_by")) if c.get(key)]
    if sign:
        add_text(s, 60.5, y, 560, 18, [P([R("      ".join(sign), 12, False, GREY)])], wrap=False, name="Sign-off")
        y += 24.0
    if c.get("ref"):
        add_text(s, 60.5, y, 560, 18, [P([R(f"Ref.: {c['ref']}", 12, False, GREY)])], wrap=False, name="Reference")
    return s


def build_material(prs, m):
    s = content_slide(prs, m.get("title", "1. Material Properties"))
    hdr, rows = m["header"], m["rows"]
    n = len(hdr)
    col_w = [856.8 / n] * n
    row_h = [61.2] + [61.2 if len(rows) == 1 else 40.0] * len(rows)
    add_table(s, 50.4, 111.6, col_w, row_h, [hdr] + rows, NAVY, [ROW_LT], 9, 11, BODY, "Material table", top_margin=4.3)
    yb = 111.6 + sum(row_h)
    add_text(s, 60.5, yb + 29.8, 820, 20, [P([R(m["units"], 14, False, BODY)])], wrap=False, name="Units")
    add_text(s, 60.5, yb + 73.1, 600, 20, [P([R(m.get("formula_label", "Allowable Von-Mises Stress:"), 14, True, NAVY)])], wrap=False, name="Formula label")
    add_text(s, 60.5, yb + 105.3, 220, 28, [P([R(m.get("formula", "\u03c3all = Syt / FOS"), 20, True, BODY)])], wrap=False, name="Formula")
    add_text(s, 269.3, yb + 107.0, 300, 22, [P([R(m["fos"], 15, True, BLUE)])], wrap=False, name="FOS")
    return s


def _picture_note(s, text):
    """one centred grey line (or two) near the bottom of a slide - e.g. the mesh statistics"""
    lines = balanced_lines(glue_units(text), 700.0, 13, False)
    return add_text(s, 130.0, 484.0, 700.0, len(lines) * 15.5 + 2, [P([R(line, 13, False, GREY)], align="c") for line in lines],
                    name="Picture note")


def build_picture_slide(prs, title, img, autocrop, alt, caption=None):
    s = content_slide(prs, title)
    if caption:                                                # e.g. mesh statistics: make room under the picture
        add_picture(s, img, (GEO_BOX[0], GEO_BOX[1], GEO_BOX[2], 474.0), autocrop, "c", "c", name=title, alt=alt)
        _picture_note(s, caption)
    else:
        add_picture(s, img, GEO_BOX, autocrop, "c", "c", name=title, alt=alt)
    return s


def _bc_side_plan(size, items, notes):
    """Geometry of the template layout: picture left, list right.  Shrinks the list font (14 -> 10 pt) if it would not fit."""
    slot_h = BC_IMG_BOX[3] - BC_IMG_BOX[1]
    if size is None:                               # no setup picture: the list starts at the left margin
        rect = (BC_IMG_BOX[0], BC_IMG_BOX[1], 0.0, 0.0, 1.0)
        ix, iy, iw, ih = rect[:4]
        dy = 0.0
        x0 = 54.0
    else:
        rect = fit(size, BC_IMG_BOX, "l", "t")
        if rect[3] < 0.85 * slot_h:                # wide screenshot: centre it vertically, the text column follows
            rect = fit(size, BC_IMG_BOX, "l", "c")
        ix, iy, iw, ih = rect[:4]
        dy = iy - BC_IMG_BOX[1]
        x0 = max(500.0, ix + iw + 6.0)
    right = 945.0
    y_list = 133.6 + dy
    lx, lw_max = x0 + 12.6, right - x0 - 12.6
    nx = x0 + 49.8
    for fs in (14, 13, 12, 11, 10):
        k = fs / 14.0
        lw = lw_max
        for _ in range(3):                         # list right edge must stay left of the corner triangle
            nl_items = sum(count_lines(glue_units(t), lw - 27, fs) for t in items)
            list_h = nl_items * LINE_BC * k + 8 * k * (len(items) - 1)
            lw = min(lw_max, tri_left(y_list + list_h + 4) - 6 - lx)
        yn = y_list + list_h + 30.0                # template: ~29 pt gap between list and Note
        nl, nright = 0.0, right + 8
        if notes:
            for _ in range(3):
                nl = sum(count_lines(glue_units(t), nright - nx - 27, fs) for t in notes)
                nright = min(right + 8, tri_left(yn + nl * LINE_NOTE * k + 6) - 4)
        bottom = (yn + nl * LINE_NOTE * k + 8) if notes else (y_list + list_h + 8)
        if bottom <= 505.0:
            break
    return dict(rect=rect, dy=dy, x0=x0, right=right, y_list=y_list, lx=lx, lw=lw, fs=fs, k=fs / 14.0, list_h=list_h, yn=yn, nx=nx,
                nright=nright, nl=nl, bottom=bottom, overflow=bottom > 505.0)


def _bc_stack_plan(size, items, notes):
    """Wide models: the setup picture across the full slide width, the list underneath (1 or 2 columns)."""
    right = STK_X1
    for _ in range(4):
        for cols in (1, 2):
            colw = (right - STK_X0 - 12.6 - (24.0 if cols == 2 else 0.0)) / cols
            half = (len(items) + 1) // 2
            groups = [items] if cols == 1 else [items[:half], items[half:]]
            hs = [sum(count_lines(glue_units(t), colw - 27, 14) for t in g) * LINE_BC + 8 * max(0, len(g) - 1) for g in groups]
            list_h = max(hs) if hs else 0.0
            if cols == 2 or list_h <= 110.0:
                break
        note_h = (30.0 + sum(count_lines(glue_units(t), right - STK_X0 - 49.8 - 27, 14) for t in notes) * LINE_NOTE) if notes else 0.0
        below = 8.0 + 24.0 + list_h + note_h + 6.0
        avail = max(110.0, (STK_BOT - 2.0 - below) - 86.0)
        rect = fit(size, (STK_X0, 86.0, STK_X1, 86.0 + avail), "c", "t")
        y_head = rect[1] + rect[3] + 8.0
        y_list = y_head + 28.0
        new_right = min(STK_X1, tri_left(y_list + list_h + 4) - 8.0)
        if abs(new_right - right) < 1.0:
            break
        right = new_right
    bottom = y_list + list_h + note_h
    return dict(rect=rect, cols=cols, groups=groups, colw=colw, y_head=y_head, y_list=y_list, list_h=list_h, note_h=note_h,
                bottom=bottom, overflow=bottom > STK_BOT + 4.0, fs=14, k=1.0)


def bc_layout_info(size, items, notes, layout="side"):
    """-> dict(scale=pt/px of the setup picture, fs, overflow) - used by the web page to predict readability."""
    plan = (_bc_stack_plan if layout == "stack" else _bc_side_plan)(size, items, notes or [])
    return dict(scale=plan["rect"][4], fs=plan["fs"], overflow=plan["overflow"])


def build_bc(prs, n, case, autocrop, layout="side"):
    title = f"Case {n}: {case['name']}"
    s = content_slide(prs, title, case.get("subtitle", "(Self Weight + Pressure)"))
    items, notes = case.get("bc_items") or [], case.get("notes") or []
    if case.get("bc_image"):
        buf, size = load_image(case["bc_image"], autocrop)
    else:                                           # no setup picture: text only (the caller decided that a slide is still useful)
        buf, size, layout = None, None, "side"
    if layout == "stack":
        return _build_bc_stack(s, n, case, buf, size, items, notes)
    plan = _bc_side_plan(size, items, notes)
    if buf is not None:
        place_picture(s, buf, plan["rect"], name=f"Case {n} boundary conditions", alt=f"Boundary conditions - {case['name']}")
    dy, x0, right, fs, k = plan["dy"], plan["x0"], plan["right"], plan["fs"], plan["k"]
    add_text(s, x0, 103.2 + dy, right - x0, 22, [P([R("Boundary Conditions", 16, True, NAVY)])], wrap=False, name="BC heading")
    paras = [P([R(glue_units(t), fs, False, BODY)], bullet="alphaUcPeriod", marL=27, indent=-27, line_pts=LINE_BC * k,
               after=(8 * k if j < len(items) - 1 else 0)) for j, t in enumerate(items)]
    add_text(s, plan["lx"], plan["y_list"], plan["lw"], plan["list_h"] + 8, paras, name="BC list")
    if notes:
        add_text(s, x0, plan["yn"], 60, 24, [P([R("Note:", 18, False, GREY)])], wrap=False, name="Note label")
        nparas = [P([R(glue_units(t), fs, False, GREY)], bullet="arabicPeriod", marL=27, indent=-27, line_pts=LINE_NOTE * k)
                  for t in notes]
        add_text(s, plan["nx"], plan["yn"] + 2.2, plan["nright"] - plan["nx"], plan["nl"] * LINE_NOTE * k + 8, nparas, name="Notes")
    return s


def _build_bc_stack(s, n, case, buf, size, items, notes):
    pl = _bc_stack_plan(size, items, notes)
    place_picture(s, buf, pl["rect"], name=f"Case {n} boundary conditions", alt=f"Boundary conditions - {case['name']}")
    add_text(s, STK_X0 + 6.0, pl["y_head"], 400, 22, [P([R("Boundary Conditions", 16, True, NAVY)])], wrap=False, name="BC heading")
    k0 = 0
    for ci, g in enumerate(pl["groups"]):
        x = STK_X0 + 12.6 + ci * (pl["colw"] + 24.0)
        paras = []
        for j, t in enumerate(g):
            p = P([R(glue_units(t), 14, False, BODY)], bullet="alphaUcPeriod", marL=27, indent=-27, line_pts=LINE_BC,
                  after=(8 if j < len(g) - 1 else 0))
            paras.append(p)
        tb = add_text(s, x, pl["y_list"], pl["colw"], pl["list_h"] + 8, paras, name="BC list" if ci == 0 else f"BC list {ci + 1}")
        if k0:
            for para in tb.text_frame.paragraphs:
                bu = para._p.pPr.find(qn("a:buAutoNum"))
                if bu is not None:
                    bu.set("startAt", str(k0 + 1))
        k0 += len(g)
    if notes:
        yn = pl["y_list"] + pl["list_h"] + 28.0
        add_text(s, STK_X0 + 6.0, yn, 60, 24, [P([R("Note:", 18, False, GREY)])], wrap=False, name="Note label")
        nx = STK_X0 + 6.0 + 49.8
        nparas = [P([R(glue_units(t), 14, False, GREY)], bullet="arabicPeriod", marL=27, indent=-27, line_pts=LINE_NOTE) for t in notes]
        add_text(s, nx, yn + 2.2, STK_X1 - nx - 20, pl["note_h"] + 8, nparas, name="Notes")
    return s


# wide plots (typical ANSYS screenshots): wider columns, heading + picture + caption centred as ONE block
CMP_COLS = ((54.0, 472.0), (488.0, 906.0))
CMP_CY = 292.0          # vertical centre of that block
HEAD_OFF = 22.6         # picture top - heading top   (template: 100.0 - 77.4)
CAP_GAP = 9.0           # caption top - picture bottom (template: 479.0 - 470.0)


def _compact_slots(panels, loaded):
    scales = [(c[1] - c[0]) / sz[0] for c, (_b, sz) in zip(CMP_COLS, loaded)]
    equal = abs(scales[0] - scales[1]) / max(scales) < 0.12     # similar pictures -> identical scale
    if equal:
        scales = [min(scales)] * 2
    for _ in range(300):
        dims = [(sz[0] * k, sz[1] * k) for (_b, sz), k in zip(loaded, scales)]
        hmax = max(d[1] for d in dims)
        cap_h = 0.0
        for p, c in zip(panels, CMP_COLS):
            if p.get("caption"):
                nlin = len(balanced_lines(glue_units(p["caption"]), c[1] - c[0], p["csize"], True))
                cap_h = max(cap_h, nlin * p["csize"] * 1.18)
        group = HEAD_OFF + hmax + ((CAP_GAP + cap_h) if cap_h else 0.0)
        y_head = CMP_CY - group / 2
        y_pic = y_head + HEAD_OFF
        hit = [(c[0] + c[1]) / 2 + w / 2 > tri_left(y_pic + h) - 8.0 for c, (w, h) in zip(CMP_COLS, dims)]
        if not any(hit):
            break
        if equal:                                              # shared scale -> shrink together
            scales = [k * 0.99 for k in scales]
        else:                                                  # otherwise only the picture that touches the triangle
            scales = [k * (0.99 if bad else 1.0) for k, bad in zip(scales, hit)]
    slots = []
    y_cap = y_pic + hmax + CAP_GAP
    for c, (w, h) in zip(CMP_COLS, dims):
        cx = (c[0] + c[1]) / 2
        cw = min(c[1] - c[0], 2 * (tri_left(y_cap + cap_h) - 6 - cx))     # caption must not touch the triangle either
        slots.append(dict(rect=(cx - w / 2, y_pic, w, h), hx=c, hy=y_head, cap=(cx - cw / 2, y_cap, cw, cap_h + 2.0)))
    return slots


def _side_slots(panels, loaded):
    """Two pictures side by side (the template layout) -> slot geometry for each.  Tall pictures (the template's gate plots)
    fill the template boxes; wide pictures use the compact layout."""
    rects = [fit(sz, p["box"], "c", "t") for (_b, sz), p in zip(loaded, panels)]
    sc = [r[4] for r in rects]
    if abs(sc[0] - sc[1]) / max(sc) < 0.12:                 # similar plots -> identical scale (looks balanced)
        common = min(sc)
        rects = [fit(sz, p["box"], "c", "t", max_scale=common) for (_b, sz), p in zip(loaded, panels)]
    slot_h = panels[0]["box"][3] - panels[0]["box"][1]
    if max(r[3] for r in rects) < 0.8 * slot_h:
        return _compact_slots(panels, loaded)
    slots = []
    for p, r in zip(panels, rects):
        x0, _, x1, _ = p["box"]
        cb = p.get("cbox")
        slots.append(dict(rect=r[:4], hx=(x0, x1), hy=77.4, cap=(cb[0], 479.0, cb[1] - cb[0], 44.0) if cb else None))
    return slots


def _draw_panel(s, p, buf, sl):
    """heading (optional) + picture + caption (optional)"""
    hx0, hx1 = sl["hx"]
    if p.get("heading") and not p.get("no_heading"):
        add_text(s, hx0, sl["hy"], hx1 - hx0, 24, [P([R(p["heading"], p.get("hsize", 18), True, TEAL_DK)], align="c")],
                 wrap=False, name=f"{p['heading']} heading")
    place_picture(s, buf, sl["rect"], name=p.get("name") or p["heading"], alt=p.get("alt") or p.get("heading") or "picture")
    if p.get("caption") and sl["cap"]:
        cx, cy, cw, chh = sl["cap"]
        lines = balanced_lines(glue_units(p["caption"]), cw, p["csize"], True)
        runs = []
        for k, ln_ in enumerate(lines):
            runs += (["BR"] if k else []) + [R(ln_, p["csize"], True, p["ccolor"])]
        add_text(s, cx, cy, cw, chh, [P(runs, align="c")], name=f"{p.get('name') or p['heading']} caption")


def build_two_up(s, panels, autocrop, divider_end=534.0):
    """Two pictures side by side, heading above each, optional caption below, vertical divider between.
    panels: [{heading, image, alt?, box, caption?, csize?, ccolor?, cbox?}, ...]"""
    loaded = [load_image(p["image"], autocrop) for p in panels]
    slots = _side_slots(panels, loaded)
    for p, (buf, _sz), sl in zip(panels, loaded, slots):
        _draw_panel(s, p, buf, sl)
    ln = s.shapes.add_connector(1, E(480.1), E(91.3), E(480.1), E(divider_end))      # vertical divider
    ln.line.width = Pt(2.04)
    ln.line.color.rgb = RGBColor.from_string(TEAL)
    ln.name = "Divider"
    return s


def _cap_height(p, width):
    if not p.get("caption"):
        return 0.0
    return len(balanced_lines(glue_units(p["caption"]), width, p["csize"], True)) * p["csize"] * 1.18


def _single_slot(p, size, with_heading=False):
    """One picture on its own slide: as large as the slide allows (heading optional, caption optional)."""
    hh = HEAD_OFF if with_heading else 0.0
    cap_h = _cap_height(p, STK_X1 - STK_X0)
    avail = (STK_BOT - STK_TOP) - hh - ((CAP_GAP + cap_h) if cap_h else 0.0)
    x, y, w, h, _sc = fit(size, (STK_X0, STK_TOP + hh, STK_X1, STK_TOP + hh + avail), "c", "c")
    cx = x + w / 2
    y_cap = y + h + CAP_GAP
    cw = max(220.0, min(STK_X1 - STK_X0, 2 * (tri_left(y_cap + cap_h) - 6 - cx))) if cap_h else 0.0
    return dict(rect=(x, y, w, h), hx=(STK_X0, STK_X1), hy=y - HEAD_OFF, cap=(cx - cw / 2, y_cap, cw, cap_h + 2.0) if cap_h else None)


def _stack_slots(panels, loaded):
    """Two pictures one above the other (very wide models)."""
    n = len(panels)
    gap = 12.0
    block = (STK_BOT - STK_TOP - gap * (n - 1)) / n

    def place(max_scale):
        out, y = [], STK_TOP
        for p, (_b, sz) in zip(panels, loaded):
            cap_h = _cap_height(p, STK_X1 - STK_X0)
            avail = block - HEAD_STACK - ((CAP_GAP_STACK + cap_h) if cap_h else 0.0)
            r = fit(sz, (STK_X0, y + HEAD_STACK, STK_X1, y + HEAD_STACK + avail), "c", "t", max_scale=max_scale)
            out.append((r, cap_h, y))
            y += block + gap
        return out

    placed = place(None)
    sc = [r[0][4] for r in placed]
    if n == 2 and abs(sc[0] - sc[1]) / max(sc) < 0.12:
        placed = place(min(sc))
    slots = []
    for (x, y, w, h, _s), cap_h, y0 in placed:
        cx = x + w / 2
        y_cap = y + h + CAP_GAP_STACK
        cw = max(220.0, min(STK_X1 - STK_X0, 2 * (tri_left(y_cap + cap_h) - 6 - cx))) if cap_h else 0.0
        slots.append(dict(rect=(x, y, w, h), hx=(STK_X0, STK_X1), hy=y0, cap=(cx - cw / 2, y_cap, cw, cap_h + 2.0) if cap_h else None))
    return slots


def build_stacked(s, panels, autocrop):
    loaded = [load_image(p["image"], autocrop) for p in panels]
    panels = [dict(p, hsize=16, csize=min(p.get("csize", 14), 13)) for p in panels]
    slots = _stack_slots(panels, loaded)
    for p, (buf, _sz), sl in zip(panels, loaded, slots):
        _draw_panel(s, p, buf, sl)
    if len(panels) == 2:
        ymid = STK_TOP + (STK_BOT - STK_TOP - 12.0) / 2 + 6.0
        ln = s.shapes.add_connector(1, E(STK_X0), E(ymid), E(STK_X1 - 40.0), E(ymid))
        ln.line.width = Pt(2.04)
        ln.line.color.rgb = RGBColor.from_string(TEAL)
        ln.name = "Divider"
    return s


def build_single(s, p, autocrop, with_heading=False):
    buf, size = load_image(p["image"], autocrop)
    p = dict(p, no_heading=not with_heading, csize=p.get("csize", 14), ccolor=p.get("ccolor", BODY))
    _draw_panel(s, p, buf, _single_slot(p, size, with_heading))
    return s


def panel_scales(kind, mode, sizes, captions=None):
    """pt-per-pixel scale of every picture for a layout - lets the web page predict how readable the legends will be.
    kind: 'results' (2 pictures) | 'views' (N pictures: geometry / mesh / extras);  mode: side | stack | separate"""
    caps = list(captions or [None] * len(sizes))
    loaded = [(None, tuple(sz)) for sz in sizes]
    if kind == "results":
        panels = [dict(box=RES_L_BOX, heading="Total deformation", caption=caps[0], csize=16, ccolor=NAVY, cbox=(60.0, 480.0)),
                  dict(box=RES_R_BOX, heading="Von-mises stress", caption=caps[1], csize=14, ccolor=BODY, cbox=(496.0, 852.0))]
    else:
        panels = [dict(box=(RES_L_BOX, RES_R_BOX)[i % 2], heading="v", caption=caps[i], csize=14, ccolor=BODY) for i in range(len(sizes))]
    out = [0.0] * len(sizes)
    step = 1 if mode == "separate" else 2
    for i in range(0, len(sizes), step):
        idx = list(range(i, min(i + step, len(sizes))))
        if len(idx) == 2 and mode == "side":
            slots = _side_slots([panels[j] for j in idx], [loaded[j] for j in idx])
        elif len(idx) == 2 and mode == "stack":
            slots = _stack_slots([dict(panels[j], hsize=16, csize=min(panels[j]["csize"], 13)) for j in idx], [loaded[j] for j in idx])
        else:
            slots = [_single_slot(panels[j], loaded[j][1], with_heading=False) for j in idx]
        for j, sl in zip(idx, slots):
            out[j] = sl["rect"][2] / loaded[j][1][0]
    return out


def image_size(spec, autocrop=True):
    """(w, h, k): size in px of the picture exactly as it will be placed (margins trimmed, huge pictures scaled down),
    and the scale k that was applied to the original pixels."""
    key = _cache_key(spec, autocrop, 0.012)
    if key is not None and key in _IMG_CACHE:
        _d, size, k, _p = _IMG_CACHE[key]
        return size[0], size[1], k
    load_image(spec, autocrop)                                  # fills the cache
    if key is not None and key in _IMG_CACHE:
        _d, size, k, _p = _IMG_CACHE[key]
        return size[0], size[1], k
    im, k = _prepare_image(spec, autocrop)
    return im.size[0], im.size[1], k


def build_results(prs, title, case, autocrop, layout="side", pre=""):
    panels = [
        dict(box=RES_L_BOX, heading="Total deformation", image=case.get("deformation_image"), caption=case.get("deformation_caption"),
             csize=16, ccolor=NAVY, cbox=(60.0, 480.0), alt="Total deformation contour plot"),
        dict(box=RES_R_BOX, heading="Von-mises stress", image=case.get("stress_image"), caption=case.get("stress_caption"),
             csize=14, ccolor=BODY, cbox=(496.0, 852.0), alt="Von-mises stress contour plot"),
    ]
    have = [p for p in panels if p["image"]]
    if not have:
        return []                                               # nothing to show: no results slide
    if len(have) == 1:                                          # only one of the two plots exists: it gets the whole slide
        p = have[0]
        ttl = f"{pre}Total Deformation" if p is panels[0] else f"{pre}Von-Mises Stress"
        sl = content_slide(prs, ttl)
        build_single(sl, p, autocrop, with_heading=False)
        return [sl]
    if layout == "separate":                                    # one big picture per slide
        out = []
        for p, ttl in zip(panels, (f"{pre}Total Deformation", f"{pre}Von-Mises Stress")):
            sl = content_slide(prs, ttl)
            build_single(sl, p, autocrop, with_heading=False)
            out.append(sl)
        return out
    s = content_slide(prs, title)
    (build_stacked if layout == "stack" else build_two_up)(s, panels, autocrop)
    return [s]


def build_view_slides(prs, title, views, autocrop, what, layout="side", single_title=None, note=None):
    """N pictures with headings (geometry views, mesh views, section / detail views).
    side: two per slide;  stack: two per slide, one above the other;  separate: one per slide.
    views: [{heading, image, caption?}]"""
    slides = []
    n = len(views)
    step = 1 if layout == "separate" else 2
    for i in range(0, n, step):
        chunk = views[i:i + step]
        if len(chunk) == 2:
            s = content_slide(prs, title if i == 0 else f"{title} (cont.)")
            panels = [dict(box=b, heading=v["heading"], image=v["image"], alt=f"{what} - {v['heading']}", caption=v.get("caption"),
                           csize=14, ccolor=BODY, cbox=cb) for v, b, cb in zip(chunk, (RES_L_BOX, RES_R_BOX), ((60.0, 480.0), (496.0, 852.0)))]
            if layout == "stack":
                build_stacked(s, panels, autocrop)
            else:
                build_two_up(s, panels, autocrop, divider_end=(476.0 if (note and i == 0) else 534.0))   # keep the note clear of the divider
            if note and i == 0:
                _picture_note(s, note)
            slides.append(s)
            continue
        v = chunk[0]
        if n == 1 and not v.get("caption") and not single_title:        # plain single picture: the template's geometry slide
            slides.append(build_picture_slide(prs, title, v["image"], autocrop, what, caption=note))
            continue
        if note and i == 0 and not v.get("caption"):
            v = dict(v, caption=note)
        ttl = single_title(v) if single_title else (title if n == 1 else f"{title} \u2013 {v['heading']}")
        s = content_slide(prs, ttl)
        build_single(s, dict(heading=v["heading"], image=v["image"], alt=f"{what} - {v['heading']}", caption=v.get("caption"),
                             csize=14, ccolor=BODY, name=v["heading"]), autocrop, with_heading=False)
        slides.append(s)
    return slides


def build_views_slide(prs, title, views, autocrop, what):
    """Geometry slide with 2 views (e.g. top / bottom), same two-up design as the results slide."""
    if len(views) != 2:
        raise ValueError("geometry.images must hold exactly 2 views here (use build_view_slides for other counts)")
    return build_view_slides(prs, title, views, autocrop, what, "side")[0]


EXTRA_LABEL = {"bc": "Boundary Conditions", "deformation": "Total Deformation", "stress": "Von-Mises Stress", "other": "Results"}


def _build_extras(prs, pre, case, kind, autocrop, layout):
    """Section / detail / additional views of one load case (after the main slide of that kind)."""
    views = [v for v in case.get("extras", []) if v.get("kind") == kind]
    if not views:
        return
    base = f"{pre}{EXTRA_LABEL[kind]}"
    if kind == "other":
        single = lambda v: f"{pre}{v['heading']}"                                         # noqa: E731
        multi = f"{pre}Additional Results"
    else:
        single = lambda v: f"{base} \u2013 {v['heading']}"                              # noqa: E731
        multi = f"{base} \u2013 Additional views"
    build_view_slides(prs, multi, views, autocrop, f"{EXTRA_LABEL[kind]} - additional view", layout, single_title=single)


def _bullet_fit(texts, width=800.0, extra_last=None, top=98.4, limit=498.0):
    """largest bullet font (17 -> 13 pt) for which the list ends above `limit` -> (size, pitch, space_after)"""
    for size in (17, 16, 15, 14, 13):
        k = size / 17.0
        total = 0.0
        for j, t in enumerate(texts):
            n = count_lines(glue_units(t), width - 12, size)
            if extra_last and j == len(texts) - 1:
                n += count_lines(extra_last, width - 12, size)
            total += n * LINE_OBS * k + 8 * k
        if top + total <= limit:
            break
    return size, LINE_OBS * size / 17.0, 8 * size / 17.0


def build_observation(prs, title, bullets, conclusion):
    if not bullets and not conclusion:
        return None
    s = content_slide(prs, title)
    size, pitch, after = _bullet_fit(bullets, extra_last=conclusion)
    paras = []
    for k, t in enumerate(bullets):
        runs = [R(glue_units(t), size, False, BODY)]
        if k == len(bullets) - 1 and conclusion:
            runs += ["BR", R(conclusion, size, False, BODY)]
        paras.append(P(runs, bullet="char", marL=12, indent=-12, line_pts=pitch, after=after))
    add_text(s, 60.1, 98.4, 800, 200, paras, name="Observations")
    return s


def build_assumptions(prs, title, bullets):
    """Assumptions & scope: a plain bullet slide (only built when the engineer ticked / typed at least one line)."""
    s = content_slide(prs, title)
    size, pitch, after = _bullet_fit(bullets)
    paras = [P([R(glue_units(t), size, False, BODY)], bullet="char", marL=12, indent=-12, line_pts=pitch, after=after) for t in bullets]
    add_text(s, 60.1, 98.4, 800, 200, paras, name="Assumptions")
    return s


def build_summary(prs, sm):
    hdr, rows = sm["header"], sm["rows"]
    per = int(sm.get("rows_per_slide", SUMMARY_ROWS))
    chunks = [rows[i:i + per] for i in range(0, len(rows), per)] or [[]]
    out = []
    for k, ch in enumerate(chunks):
        s = content_slide(prs, sm.get("title", "Final Summary") + (" (cont.)" if k else ""))
        col_w = [798.0 / len(hdr)] * len(hdr)
        add_table(s, 54.0, 108.0, col_w, [37.5] * (len(ch) + 1), [hdr] + ch, TEAL, [BAND1, BAND2], 11, 13, TXT_GREY, "Summary table", top_margin=4.5)
        out.append(s)
    return out[0]


# ─────────────────────────────── derived text (auto wording) ───────────────────────────────
_MM_MODE = {"mode": "sample"}


def fmt_mm(v: float) -> str:
    """'sample' (default): whole mm when >= 1 else 2 decimals (as in the template).
    'sig3': keep the legend's own 3 significant digits (0.085 -> 0.085, 2.42 -> 2.42)."""
    if _MM_MODE["mode"] == "sig3":
        return f"{float(f'{v:.3g}'):g}"
    return f"{round(v):d}" if v >= 1 else f"{v:.2f}"


def fmt_mpa(v: float) -> str:
    return f"{v:,.0f}".replace(",", "") if v >= 100 else f"{v:.1f}"


def derive_case_texts(cfg, case):
    """Fill in captions / observations from the max values when not given explicitly."""
    obj = cfg.get("object_short", "GG")
    allow = cfg.get("allowable_mpa")
    d, sx = case.get("max_deformation_mm"), case.get("max_stress_mpa")
    exceeds = bool(case.get("stress_exceeds")) if "stress_exceeds" in case else (sx is not None and allow is not None and sx > allow)
    case.setdefault("_exceeds", exceeds)
    if d is not None:
        case.setdefault("deformation_caption", f"Maximum Total Deformation \u2013 {fmt_mm(d)} mm")
    if allow is not None:
        case.setdefault("stress_caption",
                        f"The stress exceeds the materials allowable limit of {allow:g} MPa." if exceeds
                        else f"The stress is within the materials allowable limit of {allow:g} MPa.")
    if "observations" not in case and d is not None and sx is not None and allow is not None:
        if exceeds:
            case["observations"] = [
                f"Total deformation in the {obj} is {fmt_mm(d)} mm.",
                f"Von-Mises stress in the {obj} exceeds the materials allowable stress limit of {allow:g} MPa.",
                f"Therefore, the {obj} is structurally inadequate to withstand the applied load under this load case."]
            case.setdefault("conclusion", "Need to take corrective actions.")
        else:
            case["observations"] = [
                f"Total deformation in the {obj} is {fmt_mm(d)} mm.",
                f"Von-Mises stress in the {obj} ({fmt_mpa(sx)} MPa) is within the materials allowable stress limit of {allow:g} MPa.",
                f"Therefore, the {obj} is structurally adequate to withstand the applied load under this load case."]
            case.setdefault("conclusion", None)


def derive_summary(cfg):
    sm = cfg.setdefault("summary", {})
    sm.setdefault("header", ["Load Case", "Max. Deformation", "Max. Von-Mises Stress", "Allowable Stress"])
    if "rows" not in sm:
        sm["rows"] = []
        allow = cfg.get("allowable_mpa")
        for c in cfg["cases"]:
            sx, d = c.get("max_stress_mpa"), c.get("max_deformation_mm")
            stress = "\u2013" if sx is None else ("Exceeding Limit" if c["_exceeds"] else f"{fmt_mpa(sx)} MPa")
            sm["rows"].append([f"Case \u2013 {c.get('short_name', c['name'])}", "\u2013" if d is None else f"{fmt_mm(d)} mm",
                               stress, "[not entered]" if allow is None else f"{allow:g} MPa"])
    return sm


# ───────────────────────────────────────── main ─────────────────────────────────────────
def _has_pics(v) -> bool:
    return bool(v) and bool(v.get("image") or v.get("images"))


def build_report(cfg: dict, out: str | None = None) -> str:
    """cfg["missing"]: what to do with a picture that does not exist - "skip" (default: the slide is left out) or
    "frame" (an empty labelled frame is kept, to be filled in PowerPoint with Change Picture).  What was done is listed
    in cfg["_notes"].  A missing picture or number never raises."""
    ac = cfg.get("autocrop", True)
    _MM_MODE["mode"] = cfg.get("mm_format", "sample")
    LAY = cfg.get("layouts") or {}                              # {"results"|"geometry"|"bc"|"extras": side|stack|separate}
    policy = "frame" if cfg.get("missing") == "frame" else "skip"
    notes = cfg["_notes"] = []
    prs = build_template()
    build_cover(prs, cfg["cover"])
    build_material(prs, cfg["material"])
    obj = cfg.get("object_short", "GG")
    label = cfg.get("object_label") or (obj if obj.isupper() else obj[:1].upper() + obj[1:])
    g, m = cfg.get("geometry") or {}, cfg.get("mesh") or {}
    lay_views = LAY.get("geometry", "side")
    sec = 2                                                     # section numbers follow what is really in the report
    if _has_pics(g) or policy == "frame":
        title = g.get("title", f"{sec}. {label} Geometry")
        if g.get("images"):
            build_view_slides(prs, title, g["images"], ac, f"Geometry of the {obj}", lay_views)
        else:
            img = g.get("image") or placeholder_picture("Geometry picture not provided", size=(1200, 700))
            build_picture_slide(prs, title, img, ac, f"Geometry of the {obj}")
        sec += 1
        if not _has_pics(g):
            notes.append("Geometry: no picture - an empty frame was kept on the geometry slide.")
    else:
        notes.append("Geometry slide left out: no geometry picture.")
    stats = m.get("stats")
    if _has_pics(m) or policy == "frame" or stats:
        title = m.get("title", f"{sec}. {label} Mesh")
        if m.get("images"):
            build_view_slides(prs, title, m["images"], ac, f"Finite element mesh of the {obj}", lay_views, note=stats)
        else:
            img = m.get("image") or placeholder_picture("Mesh picture not provided", size=(1200, 700))
            build_picture_slide(prs, title, img, ac, f"Finite element mesh of the {obj}", caption=stats)
        sec += 1
        if not _has_pics(m):
            notes.append("Mesh: no picture - an empty frame was kept on the mesh slide" + (" (it carries the mesh statistics)." if stats else "."))
    else:
        notes.append("Mesh slide left out: no mesh picture.")
    if cfg.get("assumptions"):
        build_assumptions(prs, cfg.get("assumptions_title", f"{sec}. Assumptions & Scope"), cfg["assumptions"])
    cases = cfg.get("cases") or []
    multi = len(cases) > 1
    for n, case in enumerate(cases, 1):
        derive_case_texts(cfg, case)
        pre = f"Case {n}: " if multi else ""
        lay_ext = LAY.get("extras", "side")
        case.setdefault("name", f"Load Case {n}")
        has_list = bool(case.get("bc_items")) or bool(case.get("notes"))
        if policy == "frame":
            for key, ttl, sz in (("bc_image", "Setup picture not provided", (1080, 790)),
                                 ("deformation_image", "Total deformation picture not provided", (1000, 580)),
                                 ("stress_image", "Von-Mises stress picture not provided", (1000, 580))):
                if not case.get(key):
                    case[key] = placeholder_picture(ttl, size=sz)
                    notes.append(f"Case {n}: {ttl.replace(' picture not provided', '')} picture missing - an empty frame was kept.")
        if case.get("bc_image") or has_list:
            build_bc(prs, n, case, ac, case.get("bc_layout") or LAY.get("bc", "side"))
            if not case.get("bc_image"):
                notes.append(f"Case {n}: no setup picture - the setup slide holds the list of loads and supports only.")
        else:
            notes.append(f"Case {n}: setup slide left out (no setup picture and no list of loads and supports).")
        _build_extras(prs, pre, case, "bc", ac, lay_ext)
        made = build_results(prs, case.get("results_title", f"{pre}Results"), case, ac, case.get("layout") or LAY.get("results", "side"), pre)
        if not made:
            notes.append(f"Case {n}: results slide left out (no deformation and no stress picture).")
        elif not (case.get("deformation_image") and case.get("stress_image")):
            notes.append(f"Case {n}: only one result picture - it fills the slide.")
        for kind in ("deformation", "stress", "other"):
            _build_extras(prs, pre, case, kind, ac, lay_ext)
        if not build_observation(prs, case.get("observation_title", f"{pre}Observation & Summary"),
                                 case.get("observations") or [], case.get("conclusion")):
            notes.append(f"Case {n}: observation slide left out (no numbers to write about).")
    if cases:
        build_summary(prs, derive_summary(cfg))
    else:
        notes.append("No load cases: the report holds the cover, material, geometry and mesh slides only.")

    cp = prs.core_properties                                    # python-pptx refuses document properties longer than 255 characters
    short = lambda t: re.sub(r"\s+", " ", str(t or "")).strip()[:250]
    cp.title = short(cfg["cover"].get("title")) or "FEA report"
    cp.subject = short(f"FEA report {cfg['cover'].get('report_no', '')}")
    cp.author = short(cfg.get("author", "Mech Well")) or "Mech Well"
    out = out or cfg.get("output", "FEA_Report.pptx")
    Path(out).parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    return out


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    cfg_path = Path(sys.argv[1]).resolve()
    cfg = json.loads(cfg_path.read_text(encoding="utf-8"))
    print("saved:", build_report(cfg))
